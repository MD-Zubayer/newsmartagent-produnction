from io import BytesIO
from collections import Counter

from django.conf import settings
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from rest_framework import status
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView

from .auth import active_collection
from .correction import FIELDS, correct_entries
from .models import Collection, District, Entry, PromptConfig
from .admin_auth import AssunnahAdminPermission
from .prompt_service import DEFAULT_PROMPTS, effective_prompt, extract_auto_fields, resolve_district
from .serializers import DistrictSerializer, EntrySerializer, normalize_digits, to_arabic_digits


def dakhila_key(item):
    return item["dakhila_no"], item.get("entry_language", "bn")


def find_arabic_source(collection, dakhila_no):
    """Find Arabic entries even when older rows contain Arabic/Bengali digits."""
    normalized = normalize_digits(dakhila_no)
    if not normalized:
        return None
    entries = collection.entries.select_related("district").filter(entry_language="ar")
    return next((entry for entry in entries if normalize_digits(entry.dakhila_no) == normalized), None)


def sort_entries_like_ui(entries, language=None, count_entries=None):
    """Apply the same district-frequency ordering used by the As-Sunnah UI."""
    entries = list(entries)
    count_entries = entries if count_entries is None else list(count_entries)
    district_counts = Counter(entry.district_id for entry in count_entries)

    def sort_key(entry):
        label = entry.district.name if language == "ar" or (language is None and entry.entry_language == "ar") else entry.district.display_name
        dakhila = normalize_digits(entry.dakhila_no)
        try:
            dakhila_key = (0, int(dakhila))
        except (TypeError, ValueError):
            dakhila_key = (1, str(entry.dakhila_no))
        return (-district_counts[entry.district_id], label, dakhila_key, entry.id)

    return sorted(entries, key=sort_key)


class BaseView(APIView):
    permission_classes = [AllowAny]

    def collection(self, request):
        return active_collection(request)


class PromptConfigView(BaseView):
    permission_classes = [AssunnahAdminPermission]

    def get(self, request):
        collection = self.collection(request)
        output = {}
        for feature in DEFAULT_PROMPTS:
            config = PromptConfig.objects.filter(collection=collection, feature=feature).first()
            prompt, default, mode = effective_prompt(collection, feature)
            output[feature] = {"default_prompt": default, "effective_prompt": prompt, "custom_prompt": config.custom_prompt if config else "", "mode": config.mode if config else "default"}
        return Response(output)

    def put(self, request):
        collection = self.collection(request)
        feature = str(request.data.get("feature", ""))
        if feature not in DEFAULT_PROMPTS:
            return Response({"detail": "Invalid AI feature"}, status=400)
        mode = str(request.data.get("mode", "append"))
        if mode not in {"default", "append", "replace"}:
            return Response({"detail": "Invalid prompt mode"}, status=400)
        custom_prompt = str(request.data.get("custom_prompt", ""))[:12000]
        config, _ = PromptConfig.objects.update_or_create(collection=collection, feature=feature, defaults={"mode": mode, "custom_prompt": custom_prompt})
        prompt, default, _ = effective_prompt(collection, feature)
        return Response({"default_prompt": default, "effective_prompt": prompt, "custom_prompt": config.custom_prompt, "mode": config.mode})


class AutoFieldFillView(BaseView):
    def post(self, request):
        raw_text = str(request.data.get("raw_text", "")).strip()
        if not raw_text:
            return Response({"detail": "কাঁচা data দিন"}, status=400)
        if len(raw_text) > 12000:
            return Response({"detail": "একবারে সর্বোচ্চ 12000 character দেওয়া যাবে"}, status=400)
        collection = self.collection(request)
        districts = list(collection.districts.values_list("name", flat=True))
        try:
            records = extract_auto_fields(collection, raw_text, districts)
        except Exception as exc:
            return Response({"detail": f"AI Auto Field Fill ব্যর্থ হয়েছে: {exc}"}, status=502)
        for record in records:
            source = find_arabic_source(collection, record.get("dakhila_no", "")) if record.get("dakhila_no") else None
            if source:
                record["mobile_number"] = source.mobile_number
                record["blood_group"] = source.blood_group
                record["education_period"] = source.education_period
                record["district_name"] = source.district.name
        return Response({"records": records, "model": getattr(settings, "ASSUNNAH_AI_MODEL", "gpt-4o")})


class AutoFieldFillApproveView(BaseView):
    def post(self, request):
        collection = self.collection(request)
        rows = request.data.get("entries", [])
        if not rows:
            return Response({"detail": "Approve করার জন্য কোনো row নেই"}, status=400)
        prepared_rows = []
        for raw_row in rows:
            row = dict(raw_row)
            row["entry_language"] = row.get("entry_language") or "bn"
            if row["entry_language"] == "bn" and row.get("dakhila_no"):
                source = find_arabic_source(collection, row["dakhila_no"])
                if source:
                    row["mobile_number"] = source.mobile_number
                    row["blood_group"] = source.blood_group
                    row["education_period"] = source.education_period
                    row.setdefault("district", source.district_id)
            district_name = str(row.pop("district_name", "") or "").strip()
            if not row.get("district") and district_name:
                district = collection.districts.filter(name__iexact=district_name).first()
                if not district:
                    district = District.objects.create(collection=collection, name=district_name[:100])
                row["district"] = district.pk
            prepared_rows.append(row)
        serializer = EntrySerializer(data=prepared_rows, many=True)
        serializer.is_valid(raise_exception=True)
        keys = [dakhila_key(item) for item in serializer.validated_data]
        existing = set(collection.entries.filter(dakhila_no__in=[key[0] for key in keys], entry_language__in=[key[1] for key in keys]).values_list("dakhila_no", "entry_language"))
        existing = existing.intersection(keys)
        if existing:
            return Response({"detail": "একই ভাষায় এই দাখিলা নম্বর আগে থেকেই entry করা আছে"}, status=400)
        districts = {item["district"].pk for item in serializer.validated_data}
        valid = set(collection.districts.filter(pk__in=districts).values_list("pk", flat=True))
        if districts != valid:
            return Response({"detail": "এক বা একাধিক district সঠিক নয়"}, status=400)
        with transaction.atomic():
            entries = [Entry(collection=collection, **item) for item in serializer.validated_data]
            Entry.objects.bulk_create(entries)
        return Response(EntrySerializer(entries, many=True).data, status=status.HTTP_201_CREATED)


class ArabicSourceView(BaseView):
    def get(self, request):
        dakhila_no = normalize_digits(request.query_params.get("dakhila_no", ""))
        if not dakhila_no:
            return Response({"detail": "দাখিলা নম্বর দিন"}, status=400)
        entry = find_arabic_source(self.collection(request), dakhila_no)
        if not entry:
            return Response({"detail": "এই দাখিলার Arabic entry পাওয়া যায়নি"}, status=404)
        return Response({
            "id": entry.id,
            "dakhila_no": entry.dakhila_no,
            "mobile_number": entry.mobile_number,
            "blood_group": entry.blood_group,
            "education_period": entry.education_period,
            "district": DistrictSerializer(entry.district).data,
        })


class DistrictResolveView(BaseView):
    def post(self, request):
        collection = self.collection(request)
        fields = request.data if isinstance(request.data, dict) else {}
        if not any(str(fields.get(field, "")).strip() for field in ("name", "father_name", "village", "thana")):
            return Response({"detail": "District resolve করার জন্য entry fields দিন"}, status=400)
        districts = list(collection.districts.values_list("name", flat=True))
        try:
            district_name = resolve_district(collection, fields, districts)
        except Exception as exc:
            return Response({"detail": f"AI district resolve ব্যর্থ হয়েছে: {exc}"}, status=502)
        if not district_name:
            return Response({"district": None, "created": False})
        district = collection.districts.filter(name__iexact=district_name).first()
        created = False
        if not district:
            district = District.objects.create(collection=collection, name=district_name[:100])
            created = True
        return Response({"district": DistrictSerializer(district).data, "created": created})


class StateView(BaseView):
    def get(self, request):
        collection = self.collection(request)
        return Response({"authenticated": True, "expires_at": collection.expires_at,
                         "entry_count": collection.entries.count()})


class DistrictView(BaseView):
    def get(self, request):
        collection = self.collection(request)
        return Response(DistrictSerializer(collection.districts.all(), many=True).data)

    def post(self, request):
        collection = self.collection(request)
        name = str(request.data.get("name", "")).strip()
        if not name:
            return Response({"detail": "জেলার নাম দিন"}, status=400)
        district, _ = District.objects.get_or_create(collection=collection, name=name)
        return Response(DistrictSerializer(district).data, status=status.HTTP_201_CREATED)


class DistrictDetailView(BaseView):
    permission_classes = [AssunnahAdminPermission]
    def patch(self, request, pk):
        collection = self.collection(request)
        district = collection.districts.filter(pk=pk).first()
        if not district:
            return Response({"detail": "District পাওয়া যায়নি"}, status=404)
        name = str(request.data.get("name", district.name)).strip()
        name_bn = str(request.data.get("name_bn", district.name_bn)).strip()
        if not name:
            return Response({"detail": "Arabic district name দিন"}, status=400)
        if collection.districts.filter(name__iexact=name).exclude(pk=district.pk).exists():
            return Response({"detail": "এই district আগে থেকেই আছে"}, status=400)
        district.name = name[:100]
        district.name_bn = name_bn[:100]
        district.save(update_fields=["name", "name_bn"])
        return Response(DistrictSerializer(district).data)

    def delete(self, request, pk):
        collection = self.collection(request)
        district = collection.districts.filter(pk=pk).first()
        if not district:
            return Response({"detail": "District পাওয়া যায়নি"}, status=404)
        replacement_id = request.data.get("replacement_district")
        entries = district.entries.all()
        if entries.exists():
            replacement = collection.districts.filter(pk=replacement_id).exclude(pk=district.pk).first()
            if not replacement:
                return Response({"detail": "এই district-এ entry আছে। আগে target district নির্বাচন করুন।"}, status=400)
            entries.update(district=replacement)
        district.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class DistrictMapView(BaseView):
    permission_classes = [AssunnahAdminPermission]

    def post(self, request, pk):
        collection = self.collection(request)
        source = collection.districts.filter(pk=pk).first()
        target = collection.districts.filter(pk=request.data.get("target_district")).first()
        if not source or not target or source.pk == target.pk:
            return Response({"detail": "Source ও target district সঠিকভাবে নির্বাচন করুন"}, status=400)
        arabic_name = source.arabic_name
        bangla_name = target.display_name
        target.name = arabic_name[:100]
        target.name_bn = bangla_name[:100]
        target.save(update_fields=["name", "name_bn"])
        moved = source.entries.update(district=target)
        source.delete()
        return Response({"merged": moved, "source_district": pk, "target_district": target.pk, "arabic_name": target.name, "bangla_name": target.name_bn})


class EntryView(BaseView):
    def get(self, request):
        collection = self.collection(request)
        entries = collection.entries.select_related("district").all()
        district = request.query_params.get("district")
        search = request.query_params.get("search", "").strip()
        if district:
            entries = entries.filter(district_id=district)
        if search:
            from django.db.models import Q
            entries = entries.filter(Q(name__icontains=search) | Q(dakhila_no__icontains=search) |
                                     Q(mobile_number__icontains=search))
        return Response(EntrySerializer(entries, many=True).data)

    def post(self, request):
        collection = self.collection(request)
        serializer = EntrySerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        district = District.objects.filter(pk=serializer.validated_data["district"].pk,
                                            collection=collection).first()
        if not district:
            return Response({"detail": "Invalid district"}, status=400)
        if Entry.objects.filter(collection=collection, dakhila_no=serializer.validated_data["dakhila_no"], entry_language=serializer.validated_data["entry_language"]).exists():
            return Response({"detail": "এই ভাষায় এই দাখিলা নম্বরের data আগে থেকেই entry করা আছে"}, status=400)
        entry = serializer.save(collection=collection)
        return Response(EntrySerializer(entry).data, status=status.HTTP_201_CREATED)


class BulkEntryView(BaseView):
    def post(self, request):
        collection = self.collection(request)
        rows = request.data if isinstance(request.data, list) else request.data.get("entries", [])
        if not rows:
            return Response({"detail": "কমপক্ষে একটি entry দিন"}, status=400)
        serializer = EntrySerializer(data=rows, many=True)
        serializer.is_valid(raise_exception=True)
        dakhila_keys = [dakhila_key(item) for item in serializer.validated_data]
        if len(dakhila_keys) != len(set(dakhila_keys)):
            return Response({"detail": "Bulk list-এ একই ভাষায় একই দাখিলা নম্বর একাধিকবার আছে"}, status=400)
        existing = set(collection.entries.filter(dakhila_no__in=[key[0] for key in dakhila_keys], entry_language__in=[key[1] for key in dakhila_keys]).values_list("dakhila_no", "entry_language"))
        existing = existing.intersection(dakhila_keys)
        if existing:
            return Response({"detail": "একই ভাষায় এই দাখিলা নম্বর আগে থেকেই entry করা আছে"}, status=400)
        district_ids = {item["district"].pk for item in serializer.validated_data}
        valid_districts = set(collection.districts.filter(pk__in=district_ids).values_list("pk", flat=True))
        if district_ids != valid_districts:
            return Response({"detail": "এক বা একাধিক district সঠিক নয়"}, status=400)
        with transaction.atomic():
            entries = [Entry(collection=collection, **item) for item in serializer.validated_data]
            Entry.objects.bulk_create(entries)
        return Response(EntrySerializer(entries, many=True).data, status=status.HTTP_201_CREATED)


class ArabicCorrectionView(BaseView):
    def post(self, request):
        collection = self.collection(request)
        requested_ids = request.data.get("entry_ids") or []
        entries = collection.entries.filter(entry_language="ar", correction_status__in=["not_checked", "rejected"])
        if requested_ids:
            entries = entries.filter(pk__in=requested_ids)
        entries = list(entries[:50])
        if not entries:
            return Response({"detail": "নতুন checking-এর জন্য কোনো Arabic entry নেই; approved ও pending suggestion skip করা হয়েছে"}, status=400)
        try:
            corrections = correct_entries(entries)
        except Exception as exc:
            return Response({"detail": f"Arabic checking ব্যর্থ হয়েছে: {exc}"}, status=502)
        output = []
        for entry in entries:
            correction = corrections[entry.id]
            original = {field: getattr(entry, field) for field in FIELDS}
            corrected = {field: correction[field] for field in FIELDS}
            entry.original_data = original
            entry.corrected_data = corrected
            entry.correction_notes = correction.get("notes", "")
            entry.correction_status = "suggested"
            entry.corrected_at = timezone.now()
            entry.save(update_fields=["original_data", "corrected_data", "correction_notes", "correction_status", "corrected_at", "updated_at"])
            output.append(entry)
        return Response(EntrySerializer(output, many=True).data)


class ApproveCorrectionView(BaseView):
    def post(self, request, pk):
        collection = self.collection(request)
        entry = collection.entries.get(pk=pk)
        if entry.entry_language != "ar" or not entry.corrected_data:
            return Response({"detail": "এই entry-এর কোনো Arabic correction নেই"}, status=400)
        for field in FIELDS:
            setattr(entry, field, entry.corrected_data.get(field, getattr(entry, field)))
        entry.correction_status = "approved"
        entry.save(update_fields=[*FIELDS, "correction_status", "updated_at"])
        return Response(EntrySerializer(entry).data)


class RejectCorrectionView(BaseView):
    def post(self, request, pk):
        collection = self.collection(request)
        entry = collection.entries.get(pk=pk)
        entry.correction_status = "rejected"
        entry.save(update_fields=["correction_status", "updated_at"])
        return Response(EntrySerializer(entry).data)


class EntryDetailView(BaseView):
    def get_object(self, request, pk):
        collection = self.collection(request)
        return Entry.objects.get(pk=pk, collection=collection)

    def patch(self, request, pk):
        entry = self.get_object(request, pk)
        serializer = EntrySerializer(entry, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        if "district" in serializer.validated_data and serializer.validated_data["district"].collection_id != entry.collection_id:
            return Response({"detail": "Invalid district"}, status=400)
        updated = serializer.save()
        if updated.entry_language == "ar" and any(field in serializer.validated_data for field in FIELDS):
            updated.correction_status = "not_checked"
            updated.original_data = {}
            updated.corrected_data = {}
            updated.correction_notes = ""
            updated.corrected_at = None
            updated.save(update_fields=["correction_status", "original_data", "corrected_data", "correction_notes", "corrected_at", "updated_at"])
        return Response(EntrySerializer(updated).data)

    def delete(self, request, pk):
        self.get_object(request, pk).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class ExportView(BaseView):
    def get(self, request):
        collection = self.collection(request)
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.shared import Pt
        document = Document()
        language = request.query_params.get("language")
        version = request.query_params.get("version", "original")
        layout = request.query_params.get("layout", "default")
        entries_query = collection.entries.select_related("district").all()
        if language in {"ar", "bn"}:
            entries_query = entries_query.filter(entry_language=language)
        entries = list(entries_query)
        # The UI counts every collection entry when ranking districts, so keep
        # that same ranking even though this document contains one language.
        all_entries = collection.entries.select_related("district").all()
        entries = sort_entries_like_ui(entries, language=language, count_entries=all_entries)
        bangla_digits = str.maketrans("0123456789", "০১২৩৪৫৬৭৮৯")
        arabic_digits = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
        if language == "ar" and layout in {"columns", "columns3"}:
            section = document.sections[0]
            section.top_margin = 457200
            section.bottom_margin = 457200
            section.left_margin = 457200
            section.right_margin = 457200
            column_count = 3 if layout == "columns3" else 2
            table = document.add_table(rows=(len(entries) + column_count - 1) // column_count, cols=column_count)
            table.style = "Table Grid"
            for index, entry in enumerate(entries):
                values = entry.corrected_data if version == "corrected" and entry.correction_status == "approved" else entry.original_data if version == "original" and entry.original_data else {}
                serial_text = str(index + 1).translate(arabic_digits)
                lines = [
                    f"{serial_text}. {values.get('name', entry.name)}",
                    f"اسم الأب: {values.get('father_name', entry.father_name)}",
                    f"القرية: {values.get('village', entry.village)}",
                    f"{values.get('thana', entry.thana)}، {entry.district.name}",
                    f"رقم الهاتف: {to_arabic_digits(entry.mobile_number)}",
                    f"فصيلة الدم: {entry.blood_group or 'غير محدد'}",
                    f"الفترة الدراسية: {to_arabic_digits(entry.education_period) + ' سنة' if entry.education_period else 'غير محدد'}",
                ]
                cell = table.cell(index // column_count, index % column_count)
                cell.text = ""
                for line_index, line in enumerate(lines):
                    paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    paragraph_properties = paragraph._p.get_or_add_pPr()
                    bidi = OxmlElement("w:bidi")
                    bidi.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "1")
                    paragraph_properties.append(bidi)
                    run = paragraph.add_run(line)
                    run.font.size = Pt(8 if layout == "columns3" else 9)
            for empty_index in range(len(entries), ((len(entries) + column_count - 1) // column_count) * column_count):
                table.cell(empty_index // column_count, empty_index % column_count).text = ""
        else:
            for serial, entry in enumerate(entries, start=1):
                values = entry.corrected_data if language == "ar" and version == "corrected" and entry.correction_status == "approved" else entry.original_data if language == "ar" and version == "original" and entry.original_data else {}
                name = values.get("name", entry.name)
                father_name = values.get("father_name", entry.father_name)
                village = values.get("village", entry.village)
                post_office = values.get("post_office", entry.post_office)
                thana = values.get("thana", entry.thana)
                if language == "ar":
                    serial_text = str(serial).translate(arabic_digits)
                    lines = [
                        f"{serial_text}. ",
                        name,
                        f"اسم الأب: {father_name}",
                        f"\u00a0{village}",
                        f"مكتب البريد: {post_office}" if post_office else "",
                        f"\u00a0{thana}، {entry.district.name}",
                        f"رقم الهاتف: {to_arabic_digits(entry.mobile_number)}",
                        f"فصيلة الدم: {entry.blood_group or 'غير محدد'}",
                        f"الفترة الدراسية: {to_arabic_digits(entry.education_period) + ' سنة' if entry.education_period else 'غير محدد'}",
                    ]
                else:
                    serial_text = str(serial).translate(bangla_digits)
                    lines = [
                        f"{serial_text}। নাম: {entry.name}",
                        f"পিতা: {entry.father_name}",
                        f"ঠিকানা: {entry.village}",
                        f"পোস্ট অফিস: {post_office}" if post_office else "",
                        f"{entry.thana}, {entry.district.display_name}",
                        f"মোবাইল: {entry.mobile_number}",
                        f"রক্তের গ্রুপ: {entry.blood_group or '—'}",
                        f"শিক্ষাকাল: {entry.education_period or '—'}",
                    ]
                for line in lines:
                    paragraph = document.add_paragraph(line)
                    if entry.entry_language == "ar":
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        paragraph_properties = paragraph._p.get_or_add_pPr()
                        bidi = OxmlElement("w:bidi")
                        bidi.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "1")
                        paragraph_properties.append(bidi)
                    else:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                document.add_paragraph()
        output = BytesIO()
        document.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        filename = "as-sunnah-arabic-columns-3.docx" if language == "ar" and layout == "columns3" else "as-sunnah-arabic-columns.docx" if language == "ar" and layout == "columns" else "as-sunnah-arabic.docx" if language == "ar" else "as-sunnah-bangla.docx" if language == "bn" else "as-sunnah.docx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
